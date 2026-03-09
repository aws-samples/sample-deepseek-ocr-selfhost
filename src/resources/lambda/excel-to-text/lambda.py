import os
import io
import boto3

s3 = boto3.client('s3', region_name=os.environ.get('REGION', 'us-east-1'))
FILES_BUCKET = os.environ.get('FILES_BUCKET', '')


def _sheet_to_markdown_xlsx(wb, sheet_name):
    """Convert an openpyxl worksheet to a markdown table."""
    ws = wb[sheet_name]
    rows = []
    for row in ws.iter_rows(values_only=True):
        # Keep rows that have at least one non-empty cell
        if any(cell is not None and str(cell).strip() for cell in row):
            rows.append(row)

    if not rows:
        return ''

    max_cols = max(len(r) for r in rows)
    lines = []

    for idx, row in enumerate(rows):
        cells = [str(cell if cell is not None else '') for cell in row]
        while len(cells) < max_cols:
            cells.append('')
        lines.append('| ' + ' | '.join(cells) + ' |')
        if idx == 0:
            lines.append('| ' + ' | '.join(['---'] * max_cols) + ' |')

    return f'Sheet: {sheet_name}\n\n' + '\n'.join(lines)


def _sheet_to_markdown_xls(wb, sheet_name):
    """Convert an xlrd worksheet to a markdown table."""
    ws = wb.sheet_by_name(sheet_name)
    rows = []
    for i in range(ws.nrows):
        row_values = ws.row_values(i)
        if any(str(v).strip() for v in row_values):
            rows.append(row_values)

    if not rows:
        return ''

    max_cols = max(len(r) for r in rows)
    lines = []

    for idx, row in enumerate(rows):
        cells = [str(cell if cell is not None else '') for cell in row]
        while len(cells) < max_cols:
            cells.append('')
        lines.append('| ' + ' | '.join(cells) + ' |')
        if idx == 0:
            lines.append('| ' + ' | '.join(['---'] * max_cols) + ' |')

    return f'Sheet: {sheet_name}\n\n' + '\n'.join(lines)


def handler(event, context):
    """Convert an Excel file (.xlsx/.xls) into markdown text per sheet."""
    s3_key = event['s3Key']
    bucket = event.get('bucket', FILES_BUCKET)

    filename = os.path.splitext(os.path.basename(s3_key))[0]

    # Download file from S3
    response = s3.get_object(Bucket=bucket, Key=s3_key)
    file_bytes = response['Body'].read()

    is_xlsx = s3_key.lower().endswith('.xlsx')
    generated = []

    if is_xlsx:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        sheet_names = wb.sheetnames

        for i, name in enumerate(sheet_names):
            markdown = _sheet_to_markdown_xlsx(wb, name)
            if markdown.strip():
                generated.append({
                    'pageNumber': i + 1,
                    'markdown': markdown,
                    'filename': filename,
                    'sheetName': name,
                })

        wb.close()
    else:
        import xlrd
        wb = xlrd.open_workbook(file_contents=file_bytes)
        sheet_names = wb.sheet_names()

        for i, name in enumerate(sheet_names):
            markdown = _sheet_to_markdown_xls(wb, name)
            if markdown.strip():
                generated.append({
                    'pageNumber': i + 1,
                    'markdown': markdown,
                    'filename': filename,
                    'sheetName': name,
                })

    total_sheets = len(generated)
    print(f'Converted {s3_key} into {total_sheets} sheets as markdown')

    return {
        'generated': generated,
        'totalPages': total_sheets,
        'originalKey': s3_key,
        'bucket': bucket,
    }
